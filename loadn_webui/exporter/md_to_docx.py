"""md → docx（python-docx）：标题层级/粗斜体/行内 code/列表/表格/fenced code/图片/引用块。

已知限制（best-effort，写入 SKILL.md 对用户明示）：
- 无语法高亮；复杂嵌套表格拍平；数学公式不转换；html 内嵌块忽略。
"""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

_MONO = "Consolas"

_INLINE_RE = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`|\[[^\]]+\]\([^)]+\))")


def _set_cn_font(run, name: str = "Microsoft YaHei") -> None:
    run.font.name = name
    run.font.element.rPr.rFonts.set(qn("w:eastAsia"), name)


def _add_inline(par, text: str) -> None:
    for tok in _INLINE_RE.split(text):
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**") and len(tok) > 4:
            r = par.add_run(tok[2:-2]); r.bold = True; _set_cn_font(r)
        elif tok.startswith("*") and tok.endswith("*") and len(tok) > 2:
            r = par.add_run(tok[1:-1]); r.italic = True; _set_cn_font(r)
        elif tok.startswith("`") and tok.endswith("`") and len(tok) > 2:
            r = par.add_run(tok[1:-1]); r.font.name = _MONO
            r.font.size = Pt(10); r.font.color.rgb = RGBColor(0x24, 0x29, 0x2E)
        elif tok.startswith("[") and "](" in tok:
            m = re.match(r"\[([^\]]+)\]\(([^)]+)\)", tok)
            if m:
                r = par.add_run(m.group(1)); r.font.color.rgb = RGBColor(0x09, 0x69, 0xDA)
                r.underline = True; _set_cn_font(r)
        else:
            r = par.add_run(tok); _set_cn_font(r)


def convert(md_text: str, dst: Path, base_dir: Path | None = None) -> Path:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Microsoft YaHei"
    style.font.size = Pt(11)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")

    lines = md_text.splitlines()
    i, n = 0, len(lines)
    while i < n:
        line = lines[i]
        s = line.strip()
        if not s:
            i += 1
            continue
        m = re.match(r"^(#{1,6})\s+(.*)$", s)
        if m:
            h = doc.add_paragraph()
            h.style = doc.styles[f"Heading {min(len(m.group(1)), 4)}"]
            _add_inline(h, m.group(2))
            i += 1
            continue
        if s.startswith("```"):
            buf = []
            i += 1
            while i < n and not lines[i].strip().startswith("```"):
                buf.append(lines[i])
                i += 1
            i += 1
            p = doc.add_paragraph()
            r = p.add_run("\n".join(buf))
            r.font.name = _MONO; r.font.size = Pt(9)
            p.paragraph_format.left_indent = Inches(0.15)
            pPr = p._p.get_or_add_pPr()  # 底纹
            from docx.oxml import OxmlElement
            shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear")
            shd.set(qn("w:fill"), "F6F8FA")
            pPr.append(shd)
            continue
        if s.startswith("|") and i + 1 < n and re.match(r"^\s*\|[\s:|-]+\|\s*$", lines[i + 1]):
            header = [c.strip() for c in s.strip("|").split("|")]
            i += 2
            rows = []
            while i < n and lines[i].strip().startswith("|"):
                rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
                i += 1
            table = doc.add_table(rows=1 + len(rows), cols=len(header))
            table.style = "Light Grid Accent 1"
            for j, cell in enumerate(header):
                _add_inline(table.rows[0].cells[j].paragraphs[0], f"**{cell}**")
            for ri, row in enumerate(rows):
                for j in range(len(header)):
                    txt = row[j] if j < len(row) else ""
                    _add_inline(table.rows[ri + 1].cells[j].paragraphs[0], txt)
            continue
        if s.startswith(">"):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25)
            r = p.add_run(s.lstrip("> ").strip()); r.italic = True
            r.font.color.rgb = RGBColor(0x57, 0x60, 0x6A); _set_cn_font(r)
            i += 1
            continue
        m = re.match(r"^\s*([-*+])\s+(.*)$", line)
        if m:
            p = doc.add_paragraph(style="List Bullet")
            _add_inline(p, m.group(2))
            i += 1
            continue
        m = re.match(r"^\s*(\d+)[.)]\s+(.*)$", line)
        if m:
            p = doc.add_paragraph(style="List Number")
            _add_inline(p, m.group(2))
            i += 1
            continue
        m = re.match(r"^!\[([^\]]*)\]\(([^)]+)\)$", s)
        if m:
            src = m.group(2)
            if not src.startswith(("http://", "https://", "data:")) and base_dir:
                p = (base_dir / src).resolve()
                if p.is_file():
                    try:
                        doc.add_picture(str(p), width=Inches(5.8))
                        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                    except Exception:  # noqa: BLE001
                        pass
            i += 1
            continue
        if s.startswith(("---", "***", "___")) and set(s) <= set("-*_ "):
            p = doc.add_paragraph()
            pPr = p._p.get_or_add_pPr()
            from docx.oxml import OxmlElement
            pBdr = OxmlElement("w:pBdr"); bottom = OxmlElement("w:bottom")
            bottom.set(qn("w:val"), "single"); bottom.set(qn("w:sz"), "6")
            bottom.set(qn("w:color"), "D8DEE4")
            pBdr.append(bottom); pPr.append(pBdr)
            i += 1
            continue
        p = doc.add_paragraph()
        _add_inline(p, s)
        i += 1
    dst.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(dst))
    return dst


def convert_file(src: Path, dst: Path | None = None) -> Path:
    return convert(src.read_text(errors="replace"),
                   dst or src.with_suffix(".docx"), base_dir=src.parent)


if __name__ == "__main__":
    import argparse
    ap = argparse.ArgumentParser(description="md → docx")
    ap.add_argument("src")
    ap.add_argument("-o", "--out")
    a = ap.parse_args()
    out = convert_file(Path(a.src), Path(a.out) if a.out else None)
    print(f"ok: {out}")
