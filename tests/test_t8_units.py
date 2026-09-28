"""T8：titlegen / skillhub / exporter(md_to_docx) 单测（46-50% 低覆盖面）。

- titlegen：_sanitize 全形态/generate_title 三门（未启用/无 key/空文本
  → None；接口异常→None 不炸；成功→清洗）/maybe_auto_title 标记语义
  （无标记不写/有标记写+清+发 session_meta）——_chat 用注入 transport
  层 stub（被测的是 generate 的门与清洗，不是 httpx）
- skillhub：search 空查询/规范层（GitHub 条目 installable、非 GitHub
  标 False、limit 截断）/catalog 目录解析/_http_get 降级 None
- md_to_docx：convert 全要素（标题/粗斜/行内 code/列表/表格/fenced/
  图片缺失降级/引用块）真文件断言
"""
from __future__ import annotations

from loadn_webui.config import CONFIG


# ================================================================ titlegen
def test_sanitize_forms():
    from loadn_webui.integrations.titlegen import _sanitize
    assert _sanitize('"标题"') == "标题"               # 剥引号
    assert _sanitize("《书名》") == "书名"
    assert _sanitize("  句号结尾。 ") == "句号结尾"
    assert _sanitize("多  空白\n换行") == "多 空白 换行"
    assert len(_sanitize("长" * 100)) == 40            # 上限 40
    assert _sanitize("") == ""
    assert _sanitize("!!尾巴标点!!") == "!!尾巴标点"   # 尾部标点剥


async def test_generate_title_gates(monkeypatch):
    from loadn_webui.integrations import titlegen as tg
    monkeypatch.setattr(CONFIG.titlegen, "api_key", "k")
    monkeypatch.setattr(CONFIG.titlegen, "enabled", False)
    assert await tg.generate_title("任务") is None      # 未启用
    monkeypatch.setattr(CONFIG.titlegen, "enabled", True)
    monkeypatch.setattr(CONFIG.titlegen, "api_key", "")
    assert await tg.generate_title("任务") is None      # 无 key
    monkeypatch.setattr(CONFIG.titlegen, "api_key", "k")
    assert await tg.generate_title("   ") is None       # 空文本


async def test_generate_title_failure_and_success(monkeypatch):
    from loadn_webui.integrations import titlegen as tg
    monkeypatch.setattr(CONFIG.titlegen, "enabled", True)
    monkeypatch.setattr(CONFIG.titlegen, "api_key", "k")

    async def boom(system, user, max_tokens=64):
        raise RuntimeError("网关不可达")

    monkeypatch.setattr(tg, "_chat", boom)
    assert await tg.generate_title("任务") is None      # 异常降级 None

    async def ok(system, user, max_tokens=64):
        return '  "修复登录超时" 。'

    monkeypatch.setattr(tg, "_chat", ok)
    assert await tg.generate_title("任务") == "修复登录超时"


async def test_maybe_auto_title_flag_semantics(client, monkeypatch):
    """无标记不写；有标记→写标题+清标记+发 session_meta。"""
    from loadn_webui import db as db_mod
    from loadn_webui.integrations import titlegen as tg
    monkeypatch.setattr(CONFIG.titlegen, "enabled", True)
    monkeypatch.setattr(CONFIG.titlegen, "api_key", "k")

    async def ok(system, user, max_tokens=64):
        return "自动起的标题"

    monkeypatch.setattr(tg, "_chat", ok)
    r = await client.post("/api/sessions", json={"title": "起名前"})
    sid = r.json()["session"]["id"]
    # 无标记：不写
    await tg.maybe_auto_title(sid, "干活")
    with db_mod.conn() as c:
        assert db_mod.kv_get(c, tg.FLAG.format(sid=sid)) is None
    # 有标记：写+清+发事件
    with db_mod.conn() as c:
        db_mod.kv_set(c, tg.FLAG.format(sid=sid), "1")
    published = []
    from loadn_webui.engine import ENGINE
    orig = ENGINE.publish

    def spy(sid_, type_, data, turn_id=None):
        published.append((type_, data))
        return orig(sid_, type_, data, turn_id=turn_id)

    monkeypatch.setattr(ENGINE, "publish", spy)
    await tg.maybe_auto_title(sid, "干活")
    with db_mod.conn() as c:
        assert db_mod.kv_get(c, tg.FLAG.format(sid=sid)) is None
        row = c.execute("SELECT title FROM sessions WHERE id=?",
                        (sid,)).fetchone()
    assert row["title"] == "自动起的标题"
    assert published and published[-1][0] == "session_meta" \
        and published[-1][1]["title"] == "自动起的标题"


# ================================================================ skillhub
def _stub_http(monkeypatch, payload):
    from loadn_webui.integrations import skillhub as sh
    monkeypatch.setattr(sh, "_http_get", lambda url, timeout=10: payload)


def test_search_empty_query(monkeypatch):
    from loadn_webui.integrations import skillhub as sh
    assert sh.search("") == {"skills": []}              # 空查询不打网


def test_search_unreachable(monkeypatch):
    from loadn_webui.integrations import skillhub as sh
    _stub_http(monkeypatch, None)
    out = sh.search("py")
    assert out.get("error") and out["skills"] == []


def test_search_normalization(monkeypatch):
    from loadn_webui.integrations import skillhub as sh
    _stub_http(monkeypatch, {"skills": [
        {"name": "a", "description": "d", "repo_url":
         "https://github.com/x/y/tree/main/skills/a", "stars": 5},
        {"name": "b", "description": "d",
         "repo_url": "https://gitlab.com/x/y"},          # 非 GitHub
        "not-a-dict",                                      # 脏数据跳过
    ]})
    out = sh.search("q")
    names = [s["name"] for s in out["skills"]]
    assert names == ["a", "b"]
    assert out["skills"][0]["installable"] is True
    assert out["skills"][1]["installable"] is False


def test_search_limit(monkeypatch):
    from loadn_webui.integrations import skillhub as sh
    _stub_http(monkeypatch, {"skills": [
        {"name": f"s{i}", "repo_url": "https://github.com/x/y"}
        for i in range(50)]})
    assert len(sh.search("q", limit=5)["skills"]) == 5


def test_catalog_parses_flat(monkeypatch):
    """jsdelivr flat 结构 → skills 目录下的 SKILL.md 路径集合。"""
    from loadn_webui.integrations import skillhub as sh
    _stub_http(monkeypatch, {"files": [
        {"name": "/README.md"},
        {"name": "/skills/pdf/SKILL.md"},
        {"name": "/skills/artifacts-builder/SKILL.md"},
        {"name": "/skills/pdf/scripts/run.py"},          # 非 SKILL.md 不进
        {"name": "/nested/deep/SKILL.md"},                # 深层也收（…/<skill>/SKILL.md 通吃）
    ]})
    out = sh.catalog()
    assert out["skills"] == ["artifacts-builder", "deep", "pdf"]


def test_catalog_unreachable(monkeypatch):
    from loadn_webui.integrations import skillhub as sh
    _stub_http(monkeypatch, None)
    assert sh.catalog().get("error")


def test_http_get_degrades(monkeypatch):
    """真 _http_get：不可达 URL → None（不炸——市场不可达是常态）。"""
    from loadn_webui.integrations.skillhub import _http_get
    assert _http_get("http://127.0.0.1:1/nope", timeout=1) is None


# ================================================================ docx
def test_docx_full_elements(tmp_path):
    from loadn_webui.exporter import md_to_docx
    md = (
        "# 大标题\n\n"
        "正文**粗体**与*斜体*与`行内code`。\n\n"
        "## 二级\n\n"
        "- 列表甲\n- 列表乙\n\n"
        "1. 有序一\n2. 有序二\n\n"
        "| 列A | 列B |\n|---|---|\n| 1 | 2 |\n\n"
        "```python\nprint('x')\n```\n\n"
        "> 引用一行\n\n"
        "![缺图](missing-img.png)\n"
    )
    dst = tmp_path / "out.docx"
    p = md_to_docx.convert(md, dst, base_dir=tmp_path)
    assert p == dst and dst.exists() and dst.stat().st_size > 1000
    # 读回断言要素（python-docx 真解析）
    from docx import Document
    doc = Document(str(dst))
    texts = [par.text for par in doc.paragraphs]
    assert any(t == "大标题" for t in texts)
    assert any("正文" in t and "粗体" in t for t in texts)
    assert any(t == "列表甲" for t in texts)
    assert any(t == "引用一行" for t in texts)
    assert any("print('x')" in t for t in texts)        # fenced code
    assert doc.tables, "表格缺失"
    assert doc.tables[0].rows[0].cells[0].text == "列A"


def test_docx_convert_file_default_dst(tmp_path):
    from loadn_webui.exporter import md_to_docx
    src = tmp_path / "doc.md"
    src.write_text("# t\n\nhello\n", encoding="utf-8")
    out = md_to_docx.convert_file(src)                  # 缺省同名 .docx
    assert out == tmp_path / "doc.docx" and out.exists()
